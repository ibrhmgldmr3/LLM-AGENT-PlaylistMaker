"""Yuklenen dokumanlardan metin cikarir.

Cikti her zaman `[(sayfa_no, metin)]`: parcalama katmani sayfa kavramini
biliyor ve alintida "s. 4" diyebiliyor. Sayfa kavrami olmayan bicimlerde
(TXT/MD/DOCX) tek bir sanal sayfa (1) donuyor -- alternatifi `None` sayfa
tasiyan ayri bir kod yolu olurdu ve kazanci yoktu.

METIN CIKMAMASI HATA DEGIL. Taranmis (goruntu) bir PDF gecerli bir dosyadir,
yalnizca metin katmani yoktur. Bu durum cagirana `EMPTY` olarak bildiriliyor ve
kullaniciya "bu dosyadan metin cikarilamadi (taranmis olabilir)" deniyor. OCR
kapsam disi ve oyle oldugu ACIKCA soyleniyor -- sessizce bos indekslemek,
kullanicinin dosyasinin arandigini sanmasi demek olurdu.
"""

from __future__ import annotations

import logging
from pathlib import Path

from src.providers.errors import ProviderPermanentError
from src.utils.text_utils import normalize_text

_log = logging.getLogger(__name__)

# Kabul edilen uzantilar. Beyaz liste: yukleme ucu bunu kullaniyor ve
# "reddedileni say" yerine "kabul edileni say" yaklasimi yeni bir bicim
# eklendiginde guvenlik kontrolunu atlamayi imkansiz kiliyor.
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md"})


def parse_document(path: Path, filename: str | None = None) -> list[tuple[int, str]]:
    """Dosyadan `(sayfa, metin)` listesi cikarir. Metin yoksa BOS liste doner.

    `filename` verilirse bicim ondan belirlenir: sunucudaki saklanan ad
    uretilmis (`uuid4().hex + uzanti`) olabilir ve kullanicinin verdigi ad
    daha guvenilir bir ipucudur. Ikisi de ayni uzantiyi tasidigi surece fark
    etmez; ayri tutulmasinin sebebi cagiranin ikisini karistirmasini onlemek.
    """
    suffix = Path(filename or path.name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ProviderPermanentError(
            f"Desteklenmeyen dosya türü: {suffix or '(uzantısız)'}. "
            f"Desteklenenler: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if suffix == ".pdf":
        pages = _parse_pdf(path)
    elif suffix == ".docx":
        pages = _parse_docx(path)
    else:
        pages = _parse_plain_text(path)

    return [(number, text) for number, text in pages if text]


def _parse_pdf(path: Path) -> list[tuple[int, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - kurulum hatasi
        raise ProviderPermanentError(
            "`pypdf` paketi kurulu değil. `pip install -r requirements.txt` çalıştırın."
        ) from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        # Bozuk/sifreli PDF: KALICI bir durum, tekrar denemek fayda etmez.
        raise ProviderPermanentError(f"PDF okunamadı: {exc}") from exc

    pages: list[tuple[int, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            # TEK bir sayfanin cikarilamamasi dosyanin tamamini dusurmemeli:
            # gomulu bir font ya da bozuk bir icerik akisi cogu zaman tek
            # sayfayi etkiliyor ve geri kalani hala degerli.
            _log.warning("PDF sayfa %s okunamadi: %s", index, path.name, exc_info=True)
            text = ""
        pages.append((index, normalize_text(text)))
    return pages


def _parse_docx(path: Path) -> list[tuple[int, str]]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - kurulum hatasi
        raise ProviderPermanentError(
            "`python-docx` paketi kurulu değil. `pip install -r requirements.txt` çalıştırın."
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ProviderPermanentError(f"DOCX okunamadı: {exc}") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    # Tablolardaki metin de aliniyor: ders notlarinda tanim/formul tablolari
    # sik ve yalnizca paragraflara bakan bir okuyucu onlari sessizce atlardi.
    for table in document.tables:
        for row in table.rows:
            blocks.extend(cell.text for cell in row.cells)

    # DOCX'te SAYFA kavrami yok (sayfalama goruntuleyicide olusuyor); tek
    # sanal sayfa donuyoruz. Paragraf numarasini sayfa gibi sunmak, alintida
    # kullaniciya YANLIS bir konum gostermek olurdu.
    return [(1, normalize_text("\n".join(block for block in blocks if block.strip())))]


def _parse_plain_text(path: Path) -> list[tuple[int, str]]:
    # `errors="replace"`: kodlamasi bozuk bir dosya yuzunden yukleme
    # dusmemeli. Bozulan karakterler arama kalitesini biraz dusurur, dosyayi
    # tamamen reddetmek ise kullaniciya hicbir sey vermez.
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise ProviderPermanentError(f"Dosya okunamadı: {exc}") from exc
    return [(1, normalize_text(raw))]
